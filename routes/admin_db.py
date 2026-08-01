from flask import Blueprint, render_template, redirect, url_for, flash, send_file
from flask_login import login_required, current_user
from models import db, InventoryItem, InventoryCategory, StockTransaction, User
from datetime import datetime
from sqlalchemy import text
import io

admin_db_bp = Blueprint('admin_db', __name__)


def admin_only(f):
    from functools import wraps
    @wraps(f)
    def decorated(*args, **kwargs):
        if not current_user.is_authenticated or current_user.role != 'admin':
            flash('Admin only.', 'danger')
            return redirect(url_for('auth.login'))
        return f(*args, **kwargs)
    return decorated


@admin_db_bp.route('/admin/db-status')
@login_required
@admin_only
def db_status():
    engine = db.engine
    with engine.connect() as conn:
        item_rows  = conn.execute(text("SELECT branch, COUNT(*) as cnt FROM inventory_items GROUP BY branch ORDER BY branch")).fetchall()
        user_rows  = conn.execute(text("SELECT branch, role, COUNT(*) as cnt FROM users GROUP BY branch, role ORDER BY branch")).fetchall()
        total      = conn.execute(text("SELECT COUNT(*) FROM inventory_items")).scalar()
        null_items = conn.execute(text("SELECT COUNT(*) FROM inventory_items WHERE branch IS NULL OR branch = 0")).scalar()

    b1_count = next((r[1] for r in item_rows if r[0] == 1), 0)
    b2_count = next((r[1] for r in item_rows if r[0] == 2), 0)

    return render_template('admin/db_status.html',
        item_rows=item_rows, user_rows=user_rows,
        total=total, null_items=null_items,
        b1_count=b1_count, b2_count=b2_count)


@admin_db_bp.route('/admin/backup')
@login_required
@admin_only
def backup():
    total_items = InventoryItem.query.count()
    total_tx    = StockTransaction.query.count()
    last_tx     = (StockTransaction.query
                   .order_by(StockTransaction.transaction_date.desc())
                   .first())
    return render_template('admin/backup.html',
        total_items=total_items, total_tx=total_tx, last_tx=last_tx)


@admin_db_bp.route('/admin/backup/download')
@login_required
@admin_only
def backup_download():
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    hdr_font = Font(bold=True, color='FFFFFF')
    hdr_fill = PatternFill('solid', fgColor='5C3D2E')

    def style_header(ws, row=1):
        for cell in ws[row]:
            cell.font = hdr_font
            cell.fill = hdr_fill

    def autosize(ws):
        for col in ws.columns:
            width = max((len(str(c.value or '')) for c in col), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(width + 2, 40)

    ws = wb.active
    ws.title = 'Inventory Items'
    ws.append(['ID', 'Name', 'Branch', 'Category', 'Weigh-In Unit', 'Storage Unit',
                'Main Storage Qty', 'Area Storage Qty', 'Minimum Stock',
                'Status', 'Created At', 'Updated At'])
    style_header(ws)
    for item in InventoryItem.query.order_by(InventoryItem.branch, InventoryItem.name).all():
        ws.append([
            item.id, item.name, item.branch, item.category.name if item.category else '',
            item.unit_type, item.storage_unit, item.main_storage_qty, item.area_storage_qty,
            item.minimum_stock, item.status_label,
            item.created_at.strftime('%Y-%m-%d %H:%M') if item.created_at else '',
            item.updated_at.strftime('%Y-%m-%d %H:%M') if item.updated_at else '',
        ])
    autosize(ws)

    ws2 = wb.create_sheet('Categories')
    ws2.append(['ID', 'Name', 'Slug'])
    style_header(ws2)
    for cat in InventoryCategory.query.order_by(InventoryCategory.name).all():
        ws2.append([cat.id, cat.name, cat.slug])
    autosize(ws2)

    ws3 = wb.create_sheet('Transactions')
    ws3.append(['ID', 'Date', 'Item', 'Branch', 'Category', 'Type', 'Quantity',
                'Unit', 'Remarks', 'User'])
    style_header(ws3)
    txs = (StockTransaction.query
           .order_by(StockTransaction.transaction_date.desc())
           .all())
    for t in txs:
        ws3.append([
            t.id, t.transaction_date.strftime('%Y-%m-%d %H:%M:%S'),
            t.item.name if t.item else '(deleted item)',
            t.item.branch if t.item else '',
            t.item.category.name if t.item and t.item.category else '',
            t.type_label, t.quantity,
            t.item.storage_unit if t.item else '',
            t.remarks or '',
            t.user.username if t.user else 'N/A',
        ])
    autosize(ws3)

    ws4 = wb.create_sheet('Users')
    ws4.append(['ID', 'Username', 'Full Name', 'Role', 'Branch', 'Active',
                'Last Login At', 'Last Login IP', 'Created At'])
    style_header(ws4)
    for u in User.query.order_by(User.id).all():
        ws4.append([
            u.id, u.username, u.full_name, u.role, u.branch, u.is_active,
            u.last_login_at.strftime('%Y-%m-%d %H:%M') if u.last_login_at else '',
            u.last_login_ip or '',
            u.created_at.strftime('%Y-%m-%d %H:%M') if u.created_at else '',
        ])
    autosize(ws4)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    fname = f'tricoffee_full_backup_{datetime.utcnow().strftime("%Y%m%d_%H%M")}.xlsx'
    return send_file(buf, as_attachment=True, download_name=fname,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
