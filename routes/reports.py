from flask import Blueprint, render_template, request, send_file, redirect, url_for, session
from flask_login import login_required, current_user
from models import InventoryItem, InventoryCategory, StockTransaction
from utils import admin_required
from datetime import datetime, timedelta, date as date_type
from sqlalchemy import func
import io

reports_bp = Blueprint('reports', __name__)


def _get_branch():
    from flask_login import current_user
    if current_user.role == 'admin':
        return session.get('admin_branch', 0)
    return current_user.branch



# ── helpers ───────────────────────────────────────────────────────────────────

def _date_range(report_type):
    today = datetime.utcnow()
    if report_type == 'weekly':
        start = (today - timedelta(days=today.weekday())).replace(
            hour=0, minute=0, second=0, microsecond=0)
    elif report_type == 'monthly':
        start = today.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    else:
        start = today.replace(hour=0, minute=0, second=0, microsecond=0)
    return start, today


def _build_data(start, end):
    branch = _get_branch()
    tx_q = (StockTransaction.query
             .join(InventoryItem)
             .filter(StockTransaction.transaction_date >= start,
                     StockTransaction.transaction_date <= end))
    if branch:
        tx_q = tx_q.filter(InventoryItem.branch == branch)
    txs = tx_q.order_by(StockTransaction.transaction_date.desc()).all()

    q = InventoryItem.query
    if branch:
        q = q.filter_by(branch=branch)
    items = q.order_by(InventoryItem.name).all()
    cats  = InventoryCategory.query.all()
    return dict(
        transactions   = txs,
        items          = items,
        categories     = cats,
        total_in       = sum(t.quantity for t in txs if t.transaction_type == 'supplier_in'),
        total_transfer = sum(t.quantity for t in txs if t.transaction_type == 'transfer_to_area'),
        total_out      = sum(t.quantity for t in txs if t.transaction_type == 'stock_out'),
        start_date     = start,
        end_date       = end,
    )


# ── standard reports ─────────────────────────────────────────────────────────

@reports_bp.route('/')
@login_required
@admin_required
def index():
    rtype      = request.args.get('type', 'daily')
    start, end = _date_range(rtype)
    data       = _build_data(start, end)
    return render_template('reports/index.html', report_type=rtype, **data)


# ── daily summary (staff + admin) ────────────────────────────────────────────

def _build_daily_summary(report_date):
    """
    Per-item daily summary.

    Columns explained:
      daily_weigh   – what the staff physically weighed/measured this morning
                      (count_open transactions).  RECORD ONLY – does not affect
                      any stored quantity.
      transfers_in  – stock moved from main storage → area cabinet today.
      used_qty      – stock consumed from the area cabinet today
                      (stock_out transactions from batch EOD).
      area_opening  – area cabinet level at START of day
                      = closing_qty + used_qty − transfers_in
      closing_qty   – current area_storage_qty  (the cabinet after today's usage)
    """
    day_start = datetime(report_date.year, report_date.month, report_date.day, 0,  0,  0)
    day_end   = datetime(report_date.year, report_date.month, report_date.day, 23, 59, 59)

    branch = _get_branch()

    tx_q = (StockTransaction.query
            .join(InventoryItem)
            .filter(StockTransaction.transaction_date >= day_start,
                    StockTransaction.transaction_date <= day_end))
    if branch:
        tx_q = tx_q.filter(InventoryItem.branch == branch)
    day_txs = tx_q.order_by(StockTransaction.transaction_date.asc()).all()

    item_q = InventoryItem.query
    if branch:
        item_q = item_q.filter_by(branch=branch)
    items = item_q.order_by(InventoryItem.name).all()
    categories = InventoryCategory.query.all()

    summary = []
    for item in items:
        itxs = [t for t in day_txs if t.item_id == item.id]

        # Daily weigh-in record (not stored, just for reporting)
        # day_txs is ordered oldest -> newest, so the last entry is the
        # most recent submission (handles resubmitted weigh-ins).
        open_txs     = [t for t in itxs if t.transaction_type == 'count_open']
        daily_weigh  = open_txs[-1].quantity        if open_txs else None
        weigh_time   = open_txs[-1].transaction_date if open_txs else None

        # Transfers into area cabinet today
        transfers_in = sum(t.quantity for t in itxs if t.transaction_type == 'transfer_to_area')

        # Items consumed from area cabinet today (EOD batch stock-out)
        used_qty     = sum(t.quantity for t in itxs if t.transaction_type == 'stock_out')

        # Back-calculate opening area cabinet level
        closing_qty  = item.area_storage_qty
        area_opening = closing_qty + used_qty - transfers_in

        has_activity = (daily_weigh is not None or transfers_in > 0 or used_qty > 0)

        summary.append(dict(
            item         = item,
            daily_weigh  = daily_weigh,
            weigh_time   = weigh_time,
            transfers_in = transfers_in,
            used_qty     = used_qty,
            area_opening = area_opening,
            closing_qty  = closing_qty,
            has_activity = has_activity,
        ))

    counters = {tx.user.full_name or tx.user.username
                for tx in day_txs
                if tx.transaction_type == 'count_open' and tx.user}

    stockers = {tx.user.full_name or tx.user.username
                for tx in day_txs
                if tx.transaction_type == 'stock_out' and tx.user}

    branch_label = {0: 'All Branches', 1: 'Tricoffee 1', 2: 'Tricoffee 2'}.get(branch, 'Tricoffee')

    return dict(
        report_date       = report_date,
        summary           = summary,
        categories        = categories,
        day_txs           = day_txs,
        counters          = counters,
        stockers          = stockers,
        branch_label      = branch_label,
        total_used        = sum(t.quantity for t in day_txs if t.transaction_type == 'stock_out'),
        total_transfers   = sum(t.quantity for t in day_txs if t.transaction_type == 'transfer_to_area'),
        items_weighed     = sum(1 for t in day_txs if t.transaction_type == 'count_open'),
        items_stocked_out = len({t.item_id for t in day_txs if t.transaction_type == 'stock_out'}),
    )


@reports_bp.route('/daily-summary')
@login_required
def daily_summary():
    raw = request.args.get('date', '')
    try:
        report_date = datetime.strptime(raw, '%Y-%m-%d').date()
    except ValueError:
        report_date = datetime.utcnow().date()

    show_all = request.args.get('show_all', '0') == '1'
    data     = _build_daily_summary(report_date)
    return render_template('reports/daily_summary.html',
                           show_all=show_all, **data)


@reports_bp.route('/daily-summary/pdf')
@login_required
def daily_summary_pdf():
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch, cm
    from reportlab.lib import colors
    from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle,
                                    Paragraph, Spacer, HRFlowable)

    raw = request.args.get('date', '')
    try:
        report_date = datetime.strptime(raw, '%Y-%m-%d').date()
    except ValueError:
        report_date = datetime.utcnow().date()

    data    = _build_daily_summary(report_date)
    summary = data['summary']
    cats    = data['categories']

    # ── only items that have activity or non-zero qty ──
    show_all = request.args.get('show_all', '0') == '1'
    if not show_all:
        summary = [r for r in summary
                   if r['has_activity'] or r['closing_qty'] > 0]

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=.5*inch, rightMargin=.5*inch,
                            topMargin=.45*inch, bottomMargin=.45*inch)

    COFFEE  = colors.HexColor('#5C3D2E')
    CREAM   = colors.HexColor('#FFF8E7')
    GREEN   = colors.HexColor('#1e8449')
    ORANGE  = colors.HexColor('#b7770d')
    RED     = colors.HexColor('#c0392b')
    GREY    = colors.HexColor('#888888')
    WHITE   = colors.white

    styles   = getSampleStyleSheet()
    title_s  = ParagraphStyle('T',  fontSize=17, fontName='Helvetica-Bold',
                               textColor=COFFEE, spaceAfter=2)
    sub_s    = ParagraphStyle('S',  fontSize=9,  fontName='Helvetica',
                               textColor=GREY,   spaceAfter=1)
    cat_s    = ParagraphStyle('C',  fontSize=10, fontName='Helvetica-Bold',
                               textColor=COFFEE, spaceBefore=10, spaceAfter=4)
    foot_s   = ParagraphStyle('F',  fontSize=8,  fontName='Helvetica',
                               textColor=GREY)

    els = []

    # ── header ──
    els.append(Paragraph(f'TRICOFFEE — {data.get("branch_label", "Tricoffee")} Daily Inventory Report', title_s))
    els.append(Paragraph(
        f'Date: {report_date.strftime("%A, %B %d, %Y")}   |   '
        f'Morning count by: {", ".join(data["counters"]) or "—"}   |   '
        f'Stock-out by: {", ".join(data["stockers"]) or "—"}   |   '
        f'Generated: {datetime.utcnow().strftime("%H:%M UTC")}', sub_s))
    els.append(HRFlowable(width='100%', thickness=1, color=COFFEE, spaceAfter=8))

    # ── summary row ──
    sum_data = [[
        'Items Weighed', 'Items Used (EOD)', 'Total Qty Used', 'Total Transferred In'
    ],[
        str(data['items_weighed']),
        str(data['items_stocked_out']),
        f'{data["total_used"]:g}',
        f'{data["total_transfers"]:g}',
    ]]
    sum_tbl = Table(sum_data, colWidths=[1.8*inch]*4)
    sum_tbl.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), COFFEE),
        ('TEXTCOLOR',  (0,0), (-1,0), WHITE),
        ('FONTNAME',   (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTNAME',   (0,1), (-1,1), 'Helvetica-Bold'),
        ('FONTSIZE',   (0,0), (-1,-1), 9),
        ('ALIGN',      (0,0), (-1,-1), 'CENTER'),
        ('VALIGN',     (0,0), (-1,-1), 'MIDDLE'),
        ('ROWBACKGROUNDS', (0,1), (-1,1), [CREAM]),
        ('BOX',        (0,0), (-1,-1), .5, colors.lightgrey),
        ('INNERGRID',  (0,0), (-1,-1), .3, colors.lightgrey),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
    ]))
    els.append(sum_tbl)
    els.append(Spacer(1, .15*inch))

    # ── per-category tables ──
    COL_W = [2.2*inch, .5*inch, .85*inch, .85*inch, .85*inch, .85*inch, .85*inch, .85*inch]
    HDR   = ['Item', 'Unit', 'Weigh-In', 'Transfer In', 'Cabinet Open', 'Used Today', 'Cabinet Close', 'Status']

    for cat in cats:
        cat_rows = [r for r in summary if r['item'].category_id == cat.id]
        if not cat_rows:
            continue

        els.append(Paragraph(cat.name.upper(), cat_s))

        tdata = [HDR]
        for row in cat_rows:
            item   = row['item']
            tdata.append([
                item.name,
                item.storage_unit or 'pcs',
                f'{row["daily_weigh"]:g}'  if row['daily_weigh']  is not None else '—',
                f'+{row["transfers_in"]:g}' if row['transfers_in'] else '—',
                f'{row["area_opening"]:g}',
                f'{row["used_qty"]:g}'     if row['used_qty']     else '—',
                f'{row["closing_qty"]:g}',
                item.status_label,
            ])

        tbl = Table(tdata, colWidths=COL_W, repeatRows=1)
        row_fills = []
        for i, row in enumerate(cat_rows, start=1):
            st = row['item'].status
            if st == 'out_of_stock':
                row_fills.append(('BACKGROUND', (0,i), (-1,i), colors.HexColor('#fdecea')))
            elif st == 'low_stock':
                row_fills.append(('BACKGROUND', (0,i), (-1,i), colors.HexColor('#fef9e7')))
            elif i % 2 == 0:
                row_fills.append(('BACKGROUND', (0,i), (-1,i), CREAM))

        tbl.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), COFFEE),
            ('TEXTCOLOR',  (0,0), (-1,0), WHITE),
            ('FONTNAME',   (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE',   (0,0), (-1,-1), 8),
            ('ALIGN',      (2,0), (-1,-1), 'CENTER'),
            ('ALIGN',      (0,0), (1,-1),  'LEFT'),
            ('VALIGN',     (0,0), (-1,-1), 'MIDDLE'),
            ('GRID',       (0,0), (-1,-1), .3, colors.lightgrey),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ] + row_fills))
        els.append(tbl)

    # ── signature footer ──
    els.append(Spacer(1, .3*inch))
    sig_data = [['Prepared by (Staff)', 'Verified by (Manager)', 'Date & Time']]
    sig_data.append([
        ', '.join(data['stockers']) or '___________________',
        '___________________',
        report_date.strftime('%B %d, %Y'),
    ])
    sig_tbl = Table(sig_data, colWidths=[2.5*inch, 2.5*inch, 2.5*inch])
    sig_tbl.setStyle(TableStyle([
        ('BACKGROUND',    (0,0), (-1,0), COFFEE),
        ('TEXTCOLOR',     (0,0), (-1,0), WHITE),
        ('FONTNAME',      (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE',      (0,0), (-1,-1), 9),
        ('ALIGN',         (0,0), (-1,-1), 'CENTER'),
        ('VALIGN',        (0,0), (-1,-1), 'MIDDLE'),
        ('GRID',          (0,0), (-1,-1), .5, colors.lightgrey),
        ('TOPPADDING',    (0,0), (-1,-1), 8),
        ('BOTTOMPADDING', (0,0), (-1,-1), 8),
        ('BACKGROUND',    (0,1), (-1,1), CREAM),
    ]))
    els.append(sig_tbl)

    doc.build(els)
    buf.seek(0)
    fname = f'tricoffee_daily_{report_date.strftime("%Y%m%d")}.pdf'
    return send_file(buf, as_attachment=True, download_name=fname,
                     mimetype='application/pdf')


# ── existing export routes ────────────────────────────────────────────────────

@reports_bp.route('/export/excel')
@login_required
@admin_required
def export_excel():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    rtype      = request.args.get('type', 'daily')
    start, end = _date_range(rtype)
    data       = _build_data(start, end)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f'Inventory {rtype.capitalize()}'

    hdr_font = Font(bold=True, color='FFFFFF')
    hdr_fill = PatternFill('solid', fgColor='5C3D2E')
    alt_fill = PatternFill('solid', fgColor='FFF8E7')
    center   = Alignment(horizontal='center')

    ws['A1'] = 'TRICOFFEE – Inventory Report'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = (f'Type: {rtype.capitalize()} | Period: '
                f'{start.strftime("%Y-%m-%d")} → {end.strftime("%Y-%m-%d %H:%M")}')
    ws['A3'] = f'Generated: {datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")}'
    ws.append([])

    ws.append(['SUMMARY'])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
    ws.append(['Stock In (Supplier)',  data['total_in']])
    ws.append(['Transfers to Area',    data['total_transfer']])
    ws.append(['Stock Out (Used)',      data['total_out']])
    ws.append([])

    inv_row = ws.max_row + 1
    ws.cell(inv_row, 1, 'CURRENT INVENTORY').font = Font(bold=True, size=12)
    hdrs = ['Item', 'Category', 'Unit', 'Main Storage', 'Area Storage', 'Min Stock', 'Status']
    ws.append(hdrs)
    for col, h in enumerate(hdrs, 1):
        cell = ws.cell(ws.max_row, col)
        cell.font, cell.fill, cell.alignment = hdr_font, hdr_fill, center

    for idx, item in enumerate(data['items']):
        ws.append([item.name, item.category.name, item.storage_unit or 'pcs',
                   item.main_storage_qty, item.area_storage_qty,
                   item.minimum_stock, item.status_label])
        if idx % 2 == 1:
            for col in range(1, 8):
                ws.cell(ws.max_row, col).fill = alt_fill

    ws.append([])
    tx_row = ws.max_row + 1
    ws.cell(tx_row, 1, 'TRANSACTIONS').font = Font(bold=True, size=12)
    txhdrs = ['Date', 'Item', 'Category', 'Type', 'Qty', 'Unit', 'Remarks', 'By']
    ws.append(txhdrs)
    for col, h in enumerate(txhdrs, 1):
        cell = ws.cell(ws.max_row, col)
        cell.font, cell.fill = hdr_font, hdr_fill

    for t in data['transactions']:
        ws.append([t.transaction_date.strftime('%Y-%m-%d %H:%M'),
                   t.item.name, t.item.category.name, t.type_label,
                   t.quantity, t.item.storage_unit or 'pcs', t.remarks or '',
                   t.user.username if t.user else 'N/A'])

    for col in ws.columns:
        w = max((len(str(cell.value or '')) for cell in col), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(w + 2, 42)

    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    fname = f'tricoffee_{rtype}_{datetime.utcnow().strftime("%Y%m%d")}.xlsx'
    return send_file(buf, as_attachment=True, download_name=fname,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@reports_bp.route('/export/pdf')
@login_required
@admin_required
def export_pdf():
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    rtype      = request.args.get('type', 'daily')
    start, end = _date_range(rtype)
    data       = _build_data(start, end)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=.5*inch, rightMargin=.5*inch,
                            topMargin=.5*inch,  bottomMargin=.5*inch)

    styles  = getSampleStyleSheet()
    coffee  = colors.HexColor('#5C3D2E')
    cream   = colors.HexColor('#FFF8E7')
    white   = colors.white
    els     = []

    title_s = ParagraphStyle('T', parent=styles['Title'],
                              textColor=coffee, fontSize=16, spaceAfter=4)
    sub_s   = ParagraphStyle('S', parent=styles['Normal'],
                              textColor=colors.grey, fontSize=9)
    h2_s    = ParagraphStyle('H', parent=styles['Heading2'],
                              textColor=coffee, fontSize=12, spaceBefore=12)

    els.append(Paragraph('TRICOFFEE – Inventory Report', title_s))
    els.append(Paragraph(
        f'Type: {rtype.capitalize()} | '
        f'Period: {start.strftime("%Y-%m-%d")} → {end.strftime("%Y-%m-%d %H:%M")}', sub_s))
    els.append(Spacer(1, .2*inch))

    els.append(Paragraph('Summary', h2_s))
    st = Table([['Metric','Value'],
                ['Stock In (Supplier)', str(data['total_in'])],
                ['Transfers to Area',   str(data['total_transfer'])],
                ['Stock Out (Used)',     str(data['total_out'])]],
               colWidths=[2.5*inch, 1.5*inch])
    st.setStyle(TableStyle([
        ('BACKGROUND', (0,0),(-1,0), coffee), ('TEXTCOLOR',(0,0),(-1,0), white),
        ('FONTNAME',   (0,0),(-1,0), 'Helvetica-Bold'), ('FONTSIZE',(0,0),(-1,-1), 9),
        ('GRID',       (0,0),(-1,-1), .5, colors.lightgrey),
        ('ROWBACKGROUNDS',(0,1),(-1,-1),[white, cream]),
    ]))
    els.append(st); els.append(Spacer(1, .2*inch))

    els.append(Paragraph('Current Inventory', h2_s))
    inv = [['Item','Category','Unit','Main','Area','Min','Status']]
    for item in data['items']:
        inv.append([item.name, item.category.name, item.storage_unit or 'pcs',
                    str(item.main_storage_qty), str(item.area_storage_qty),
                    str(item.minimum_stock), item.status_label])
    it = Table(inv, repeatRows=1,
               colWidths=[2.2*inch,1.6*inch,.6*inch,.7*inch,.7*inch,.6*inch,.9*inch])
    it.setStyle(TableStyle([
        ('BACKGROUND',(0,0),(-1,0), coffee),('TEXTCOLOR',(0,0),(-1,0), white),
        ('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('FONTSIZE',(0,0),(-1,-1), 8),
        ('GRID',(0,0),(-1,-1),.5,colors.lightgrey),
        ('ROWBACKGROUNDS',(0,1),(-1,-1),[white,cream]),
        ('VALIGN',(0,0),(-1,-1),'MIDDLE'),
    ]))
    els.append(it); els.append(Spacer(1, .2*inch))

    if data['transactions']:
        els.append(Paragraph('Transactions', h2_s))
        tx = [['Date','Item','Type','Qty','Unit','Remarks','User']]
        for t in data['transactions'][:60]:
            tx.append([t.transaction_date.strftime('%m/%d %H:%M'), t.item.name[:22],
                       t.type_label, str(t.quantity), t.item.storage_unit or 'pcs',
                       (t.remarks or '')[:28], t.user.username if t.user else 'N/A'])
        tt = Table(tx, repeatRows=1,
                   colWidths=[1.0*inch,2.0*inch,1.4*inch,.6*inch,.5*inch,2.0*inch,.8*inch])
        tt.setStyle(TableStyle([
            ('BACKGROUND',(0,0),(-1,0), coffee),('TEXTCOLOR',(0,0),(-1,0), white),
            ('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('FONTSIZE',(0,0),(-1,-1), 8),
            ('GRID',(0,0),(-1,-1),.5,colors.lightgrey),
            ('ROWBACKGROUNDS',(0,1),(-1,-1),[white,cream]),
        ]))
        els.append(tt)

    doc.build(els)
    buf.seek(0)
    fname = f'tricoffee_{rtype}_{datetime.utcnow().strftime("%Y%m%d")}.pdf'
    return send_file(buf, as_attachment=True, download_name=fname,
                     mimetype='application/pdf')
